import openpyxl, json, math, sys
import os as _os, sys as _sys
_HERE = _os.path.dirname(_os.path.abspath(__file__))
_sys.path.insert(0, _HERE)
def _here(name): return _os.path.join(_HERE, name)


def _workbook():
    """The group workbook, looked for beside this script.

    Any .xlsx in this folder carrying an A3_Inputs sheet will do, so the file
    can be renamed or reissued without editing this script.
    """
    import glob
    named = _here("MSE_July_28_stayhome_fix.xlsx")
    if _os.path.exists(named):
        return named
    for path in sorted(glob.glob(_os.path.join(_HERE, "*.xlsx"))):
        try:
            if "A3_Inputs" in openpyxl.load_workbook(path, read_only=True).sheetnames:
                return path
        except Exception:
            continue
    raise SystemExit(
        "analyse.py cannot find the group workbook.\n"
        "It needs an .xlsx file in this folder with a sheet named 'A3_Inputs',\n"
        "holding the twelve team-round rows of bundle prices. The submitted\n"
        "copy is MSE_July_28_stayhome_fix.xlsx.\n"
        "Looked in: " + _HERE)


import inputs
wb = openpyxl.load_workbook(_workbook(), data_only=True)
ws = wb["A3_Inputs"]
rows=[]
for r in range(5,17):
    team=ws.cell(r,1).value
    if not team: continue
    rows.append(dict(team=team,k=ws.cell(r,2).value,stage=ws.cell(r,3).value,city=ws.cell(r,4).value,
        Bnr=ws.cell(r,5).value,Brf=ws.cell(r,6).value,f=ws.cell(r,7).value,
        Blate=ws.cell(r,8).value,q=ws.cell(r,9).value,r=ws.cell(r,10).value))
A={t:inputs.advancement(t) for t in inputs.REACH}  # Groll/Zeileis conditional chain, per team
V=inputs.V                          # TickPick average paid, 2026 final
def evs(d,P,v=V):
    nr = P*(v-d["Bnr"]) + (1-P)*(-(1-d["r"])*d["Bnr"])
    rf = P*(v-d["Brf"]) + (1-P)*(-d["f"])
    wt = P*((1-d["q"])*(v-d["Blate"]) + d["q"]*0) + (1-P)*0
    return nr,rf,wt
print(f"{'team':6}{'k':>2} {'stage':22}{'city':16}{'P':>6}{'EV nr':>11}{'EV rf':>11}{'EV wait':>11}{'stay':>7}  {'optimal':22}{'margin':>10}{'VOI':>10}")
out=[]
for d in rows:
    P=A[d["team"]][d["k"]]
    nr,rf,wt=evs(d,P)
    rf_avail = d["Brf"]>d["Bnr"]
    alts=[("Book non-refundable",nr),("Book refundable",rf if rf_avail else None),("Wait",wt),("Stay home",0.0)]
    alts=[(n,x) for n,x in alts if x is not None]
    alts.sort(key=lambda t:-t[1])
    best,bv=alts[0]; margin=bv-alts[1][1]
    emax = P*max(V-d["Bnr"], 0.0) + (1-P)*0.0
    voi = emax-bv
    out.append(dict(d,P=P,nr=nr,rf=rf,wt=wt,best=best,bv=bv,margin=margin,voi=voi,rf_avail=rf_avail,emax=emax))
    print(f"{d['team']:6}{d['k']:>2} {d['stage']:22}{str(d['city'])[:15]:16}{P:>6.2f}{nr:>11.2f}{'' if rf_avail else '  --pruned':>11}" if False else
          f"{d['team']:6}{d['k']:>2} {d['stage']:22}{str(d['city'])[:15]:16}{P:>6.2f}{nr:>11.2f}{(f'{rf:11.2f}' if rf_avail else '     PRUNED')}{wt:>11.2f}{0:>7.2f}  {best:22}{margin:>10.2f}{voi:>10.2f}")
json.dump(out,open(_here("model_results.json"),"w"),indent=1)

d=[x for x in out if x["team"]=="Spain" and x["k"]==6][0]
print("\n--- Spain Final detail (matches Validation tab) ---")
for kk in ("Bnr","Brf","f","Blate","q","r"): print(f"  {kk:6}= {d[kk]}")
print(f"  P = {d['P']}, v = {V}")
print(f"  EV nr {d['nr']:.2f} | EV rf {d['rf']:.2f} | EV wait {d['wt']:.2f} | stay home 0.00")
print(f"  optimal: {d['best']} by {d['margin']:.2f}")
print(f"  E[max] with clairvoyance on A6 = {d['emax']:.2f}   VOI = {d['voi']:.2f}")
be = d["Brf"] + (1-d["P"])*d["f"]/d["P"]
print(f"  break-even v (refundable vs stay home) = {be:.2f}")
bnr_be = d["Bnr"] + (1-d["P"])*(1-d["r"])*d["Bnr"]/d["P"]
print(f"  break-even v (non-refundable vs stay home) = {bnr_be:.2f}")
print(f"  EV(rf)-EV(nr) = {d['rf']-d['nr']:.2f} (independent of v)")
# exponential utility
for rho in (2500,5000,10000,25000,50000):
    def ce(pairs):
        eu=sum(p*(1-math.exp(-x/rho)) for p,x in pairs)
        return -rho*math.log(1-eu) if eu<1 else float('inf')
    P=d["P"]
    c_nr=ce([(P,V-d["Bnr"]),(1-P,-(1-d["r"])*d["Bnr"])])
    c_rf=ce([(P,V-d["Brf"]),(1-P,-d["f"])])
    c_wt=ce([(P*(1-d["q"]),V-d["Blate"]),(P*d["q"],0),(1-P,0)])
    winner=max([("nr",c_nr),("rf",c_rf),("wait",c_wt),("stay",0.0)],key=lambda t:t[1])
    print(f"  rho={rho:>6}: CE nr {c_nr:10.2f}  rf {c_rf:10.2f}  wait {c_wt:10.2f}  stay 0.00 -> {winner[0]}")
